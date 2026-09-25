"""
Chargement de document multi-format — point d'entrée unique de la
première étape du pipeline (avant appariement/extraction).

Formats supportés :
    - Images (png, jpg, jpeg, tiff, bmp, webp) : OCR Tesseract (inchangé).
    - PDF : texte natif page par page quand il existe (aucune incertitude
      OCR — un PDF numérique contient déjà ses caractères) ; repli OCR
      uniquement pour les pages sans texte extractible (PDF scanné/image).
    - DOCX (Word) : paragraphes et cellules de tableau, avec des positions
      reconstituées pour rester compatible avec l'extraction ancrée par
      vecteur spatial (une cellule de tableau "clé" et sa voisine "valeur"
      à droite se comportent alors comme n'importe quelle paire clé/valeur
      détectée par OCR).
    - XLSX (Excel) : cellules positionnées sur une grille (ligne, colonne)
      — cas d'usage naturel pour l'ancrage spatial droite/dessous.

Dans tous les cas, le résultat est une liste de TextElementSchema : le
reste du pipeline (matching_service, extraction_service, ai_extraction_service,
ner_service) ne sait pas d'où viennent les éléments et n'a besoin d'aucune
modification.
"""

from __future__ import annotations

from pathlib import Path

from app.models.schemas import TextElementSchema, BBoxSchema
from app.services.preprocessing import enhance_image
from app.services.ocr_service import extract_text_elements

SUPPORTED_IMAGE_SUFFIXES = (".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp", ".webp")
SUPPORTED_SUFFIXES = SUPPORTED_IMAGE_SUFFIXES + (".pdf", ".docx", ".xlsx")

# Toutes les coordonnees produites ici (texte natif PDF, DOCX, XLSX) sont
# ramenees sur la meme echelle que les pages de PDF scannees rasterisees
# plus bas (dpi=300) et que des images scannees typiques. Sans cette
# normalisation, un document DOCX/XLSX/PDF-texte produisait des
# coordonnees des ordres de grandeur plus petites qu'une image (quelques
# dizaines de pixels contre plusieurs centaines/milliers) : S3 (score
# spatial, Definition 4.6) s'effondrait systematiquement meme quand la
# structure correspondait reellement au template, et le pipeline
# retombait a chaque fois sur le repli IA LayoutLMv3 (couteux en CPU) au
# lieu de l'extraction ancree (quasi instantanee). Ce n'est pas un
# pixel-perfect (on ne connait pas la resolution reelle d'un futur scan),
# mais ramener les echelles dans le meme ordre de grandeur est necessaire
# pour que la tolerance spatiale (SPATIAL_TOLERANCE=10px, Tableau 4.1) ait
# un sens sur ces formats.
_REFERENCE_DPI = 300
_PDF_POINTS_TO_PIXELS = _REFERENCE_DPI / 72.0  # 1 point PDF = 1/72 pouce


def load_elements(filepath: Path, lang: str = "eng") -> list[TextElementSchema]:
    """
    Charge un document quel que soit son format et retourne ses éléments
    textuels localisés, prêts pour l'appariement/extraction.
    """
    suffix = filepath.suffix.lower()

    if suffix == ".pdf":
        return _load_pdf(filepath, lang)
    if suffix == ".docx":
        return _load_docx(filepath)
    if suffix == ".xlsx":
        return _load_xlsx(filepath)
    if suffix in SUPPORTED_IMAGE_SUFFIXES:
        return _load_image(filepath, lang)

    raise ValueError(f"Format non supporte : {suffix}")


# ── Images (inchangé, juste déplacé derrière l'interface commune) ──


def _load_image(filepath: Path, lang: str) -> list[TextElementSchema]:
    import cv2

    img = cv2.imread(str(filepath))
    if img is None:
        raise ValueError(f"Impossible de charger l'image : {filepath}")

    elements = extract_text_elements(enhance_image(img), lang=lang)
    for elem in elements:
        elem.page_num = 1
    return elements


# ── PDF : texte natif, repli OCR page par page si nécessaire ──


def _load_pdf(filepath: Path, lang: str) -> list[TextElementSchema]:
    import pymupdf
    import numpy as np

    all_elements: list[TextElementSchema] = []

    with pymupdf.open(str(filepath)) as doc:
        for page_index, page in enumerate(doc, start=1):
            words = page.get_text("words")  # [(x0,y0,x1,y1,word,block,line,word_no), ...]

            if words:
                # Texte natif : pas d'incertitude OCR, confiance maximale.
                # Coordonnees en points PDF (1/72 pouce) -> pixels equivalent
                # 300 DPI, pour rester sur la meme echelle que le reste du
                # pipeline (cf. _PDF_POINTS_TO_PIXELS ci-dessus).
                for x0, y0, x1, y1, word, *_ in words:
                    word = word.strip()
                    if not word:
                        continue
                    all_elements.append(
                        TextElementSchema(
                            text=word,
                            bbox=BBoxSchema(
                                x1=int(x0 * _PDF_POINTS_TO_PIXELS),
                                y1=int(y0 * _PDF_POINTS_TO_PIXELS),
                                x2=int(round(x1 * _PDF_POINTS_TO_PIXELS)),
                                y2=int(round(y1 * _PDF_POINTS_TO_PIXELS)),
                            ),
                            confidence=1.0,
                            page_num=page_index,
                        )
                    )
                continue

            # Page sans texte extractible (probablement une page scannée) :
            # on la rasterise et on retombe sur l'OCR Tesseract existant.
            pixmap = page.get_pixmap(dpi=300)
            image = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(
                pixmap.height, pixmap.width, pixmap.n
            )
            if pixmap.n == 4:
                import cv2
                image = cv2.cvtColor(image, cv2.COLOR_RGBA2BGR)
            elif pixmap.n == 3:
                image = image[:, :, ::-1]  # RGB -> BGR

            page_elements = extract_text_elements(enhance_image(image), lang=lang)
            for elem in page_elements:
                elem.page_num = page_index
            all_elements.extend(page_elements)

    return all_elements


# ── DOCX : paragraphes + tableaux, positions reconstituées ──
# Constantes calibrees a l'echelle _REFERENCE_DPI (300 DPI) : ~1 ligne de
# texte 11-12pt (_LINE_HEIGHT), largeur usuelle d'une colonne de tableau de
# ~1.5 pouce (_CELL_WIDTH), largeur moyenne d'un caractere (_CHAR_WIDTH).

_LINE_HEIGHT = 72
_CELL_WIDTH = 450
_CHAR_WIDTH = 24


def _load_docx(filepath: Path) -> list[TextElementSchema]:
    import docx

    document = docx.Document(str(filepath))
    elements: list[TextElementSchema] = []
    row = 0

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            elements.append(
                TextElementSchema(
                    text=text,
                    bbox=BBoxSchema(
                        x1=0,
                        y1=row * _LINE_HEIGHT,
                        x2=min(len(text) * _CHAR_WIDTH, 3000),
                        y2=row * _LINE_HEIGHT + (_LINE_HEIGHT - 10),
                    ),
                    confidence=1.0,
                    page_num=1,
                )
            )
        row += 1

    for table in document.tables:
        for table_row in table.rows:
            for col_index, cell in enumerate(table_row.cells):
                text = cell.text.strip()
                if not text:
                    continue
                x1 = col_index * _CELL_WIDTH
                elements.append(
                    TextElementSchema(
                        text=text,
                        bbox=BBoxSchema(
                            x1=x1,
                            y1=row * _LINE_HEIGHT,
                            x2=x1 + (_CELL_WIDTH - 20),
                            y2=row * _LINE_HEIGHT + (_LINE_HEIGHT - 10),
                        ),
                        confidence=1.0,
                        page_num=1,
                    )
                )
            row += 1

    return elements


# ── XLSX : cellules positionnées sur une grille (ligne, colonne) ──
# Constantes calibrees a l'echelle _REFERENCE_DPI (300 DPI) : hauteur de
# ligne Excel par defaut (~15pt) et largeur de colonne par defaut.

_ROW_HEIGHT = 63
_COL_WIDTH = 200


def _load_xlsx(filepath: Path) -> list[TextElementSchema]:
    import openpyxl

    workbook = openpyxl.load_workbook(str(filepath), data_only=True)
    elements: list[TextElementSchema] = []

    for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                text = str(cell.value).strip()
                if not text:
                    continue

                x1 = (cell.column - 1) * _COL_WIDTH
                y1 = (cell.row - 1) * _ROW_HEIGHT
                elements.append(
                    TextElementSchema(
                        text=text,
                        bbox=BBoxSchema(
                            x1=x1,
                            y1=y1,
                            x2=x1 + (_COL_WIDTH - 15),
                            y2=y1 + (_ROW_HEIGHT - 5),
                        ),
                        confidence=1.0,
                        page_num=sheet_index,
                    )
                )

    return elements
